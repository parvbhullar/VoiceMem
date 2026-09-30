"""/memories loader: extract text from uploads, detect prose vs transcript, split into chunks.

Fixtures (PDF, DOCX) are generated at runtime. No network, no models::

    python tests/test_ingest_split.py
"""
import io
import json
import sys
import time
import unittest
import zipfile
from pathlib import Path

ROOT = Path(__file__).resolve().parent.parent
sys.path.insert(0, str(ROOT / "web"))

import ingest  # noqa: E402
from ingest import (CHUNK_CHARS, Chunk, detect_format, extract_text,  # noqa: E402
                    plan_chunks, split_prose, split_transcript)


def tiny_pdf(*lines: str) -> bytes:
    """One-page PDF with a real text layer. No lines -> blank page (the scanned-PDF case).
    Text must be ASCII without ( ) or backslash."""
    ops = "".join(f"({s}) Tj 0 -14 Td " for s in lines)
    stream = f"BT /F1 12 Tf 72 720 Td {ops}ET".encode() if lines else b""
    objs = [
        b"<< /Type /Catalog /Pages 2 0 R >>",
        b"<< /Type /Pages /Kids [3 0 R] /Count 1 >>",
        b"<< /Type /Page /Parent 2 0 R /MediaBox [0 0 612 792] "
        b"/Resources << /Font << /F1 4 0 R >> >> /Contents 5 0 R >>",
        b"<< /Type /Font /Subtype /Type1 /BaseFont /Helvetica "
        b"/Encoding /WinAnsiEncoding >>",
        b"<< /Length %d >>\nstream\n" % len(stream) + stream + b"\nendstream",
    ]
    out = bytearray(b"%PDF-1.4\n")
    offsets = []
    for i, body in enumerate(objs, 1):
        offsets.append(len(out))
        out += b"%d 0 obj\n" % i + body + b"\nendobj\n"
    xref = len(out)
    out += b"xref\n0 %d\n0000000000 65535 f \n" % (len(objs) + 1)
    out += b"".join(b"%010d 00000 n \n" % o for o in offsets)
    out += b"trailer\n<< /Size %d /Root 1 0 R >>\nstartxref\n%d\n%%%%EOF\n" % (
        len(objs) + 1, xref)
    return bytes(out)


def tiny_docx(*paras: str) -> bytes:
    import docx
    d = docx.Document()
    d.add_heading("Notes", level=1)
    for p in paras:
        d.add_paragraph(p)
    buf = io.BytesIO()
    d.save(buf)
    return buf.getvalue()


def sentences(n: int, word: str = "Alice") -> str:
    """n distinct 35-char sentences joined by spaces."""
    return " ".join(f"{word} said thing number {i:04d} today." for i in range(n))


class ExtractTest(unittest.TestCase):
    def test_txt_utf8_bom_stripped(self):
        self.assertEqual(extract_text("a.txt", "\ufeffcafé".encode("utf-8")), "café")

    def test_non_utf8_rejected(self):
        with self.assertRaisesRegex(ValueError, "UTF-8"):
            extract_text("a.txt", b"\xff\xfe\xfa")

    def test_unsupported_type_names_allowed(self):
        with self.assertRaises(ValueError) as cm:
            extract_text("setup.exe", b"MZ")
        for ext in (".txt", ".md", ".json", ".pdf", ".docx"):
            self.assertIn(ext, str(cm.exception))

    def test_pdf_text_layer(self):
        text = extract_text("a.pdf", tiny_pdf("Alice is vegetarian.", "She lives in Pune."))
        self.assertIn("Alice is vegetarian.", text)
        self.assertIn("She lives in Pune.", text)

    def test_blank_pdf_is_scanned(self):
        with self.assertRaisesRegex(ValueError, "scanned"):
            extract_text("scan.pdf", tiny_pdf())

    def test_garbage_pdf_rejected(self):
        with self.assertRaises(ValueError):
            extract_text("bad.pdf", b"this is not a pdf at all")

    def test_docx_headings_and_paragraphs(self):
        text = extract_text("a.docx", tiny_docx("Alice is vegetarian.", "She lives in Pune."))
        self.assertIn("# Notes", text)
        self.assertIn("Alice is vegetarian.", text)
        self.assertIn("She lives in Pune.", text)


class DetectTest(unittest.TestCase):
    def test_conversation_is_transcript(self):
        self.assertEqual(detect_format("Alice: hi\nBot: hello\nAlice: I am vegetarian"), "transcript")

    def test_one_label_line_in_prose_is_prose(self):
        text = ("Alice moved to Pune.\nShe works as a nurse.\nNote: ask about her sister.\n"
                "She likes tea.\nHer dog is called Rex.")
        self.assertEqual(detect_format(text), "prose")

    def test_key_value_notes_are_prose(self):
        self.assertEqual(detect_format("Diet: veg\nCity: Pune\nJob: nurse"), "prose")

    def test_json_turn_list_is_transcript(self):
        text = json.dumps([{"role": "user", "content": "hi"}, {"role": "assistant", "content": "yo"}])
        self.assertEqual(detect_format(text), "transcript")

    def test_two_line_exchange_is_transcript(self):
        self.assertEqual(detect_format("Alice: hi\nBot: hello"), "transcript")


class ProseTest(unittest.TestCase):
    def test_short_paragraphs_merge(self):
        self.assertEqual(split_prose("Alice is vegetarian.\n\nShe lives in Pune."),
                         [Chunk("Alice is vegetarian.\n\nShe lives in Pune.")])

    def test_long_paragraphs_split_under_limit(self):
        paras = "\n\n".join(sentences(8, w) for w in ("Alice", "Bob", "Carol"))
        chunks = split_prose(paras)
        self.assertGreaterEqual(len(chunks), 2)
        self.assertTrue(all(len(c.text) <= CHUNK_CHARS for c in chunks))

    def test_long_paragraph_cut_at_sentence_ends(self):
        para = sentences(60)
        self.assertGreater(len(para), 2000)
        chunks = split_prose(para)
        self.assertGreaterEqual(len(chunks), 3)
        for c in chunks:
            self.assertLessEqual(len(c.text), CHUNK_CHARS)
            self.assertTrue(c.text.endswith("today."), c.text[-30:])
        self.assertEqual(" ".join(c.text for c in chunks), para)

    def test_unpunctuated_sentence_cut_at_space(self):
        words = [f"w{i:03d}" for i in range(250)]           # 250 * 5 - 1 = 1249 chars
        chunks = split_prose(" ".join(words))
        self.assertGreaterEqual(len(chunks), 2)
        self.assertTrue(all(len(c.text) <= CHUNK_CHARS for c in chunks))
        self.assertEqual(" ".join(c.text for c in chunks).split(), words)   # no word cut in half

    def test_heading_prefixes_chunk(self):
        chunks = split_prose("# Diet\n\nShe is vegetarian.")
        self.assertEqual(len(chunks), 1)
        self.assertTrue(chunks[0].text.startswith("# Diet\n"))
        self.assertIn("She is vegetarian.", chunks[0].text)

    def test_new_heading_flushes(self):
        chunks = split_prose("# Diet\n\nShe is vegetarian.\n\n## Home\n\nShe lives in Pune.")
        self.assertEqual([c.text for c in chunks],
                         ["# Diet\nShe is vegetarian.", "# Home\nShe lives in Pune."])

    def test_whitespace_only_is_empty(self):
        self.assertEqual(split_prose("  \n\n \t\n"), [])


class TranscriptTest(unittest.TestCase):
    def test_default_owner_is_first_person(self):
        chunks = split_transcript("Bot: hello\nAlice: hi\nBob: hey\nAlice: I am vegetarian")
        self.assertEqual([c.speaker for c in chunks], ["user", "Bob", "user"])

    def test_self_label_beats_first_speaker(self):
        # "Me:" is the owner even when someone else speaks first; Tom keeps his name.
        chunks = split_transcript("Tom: Flights are booked.\nMe: Great.\nTom: I will handle the hotel.")
        self.assertEqual([c.speaker for c in chunks], ["Tom", "user", "Tom"])

    def test_named_owner_case_insensitive(self):
        chunks = split_transcript("Alice: hi\nBob: I live in Pune", owner="bob")
        self.assertEqual(chunks, [Chunk("hi", "Alice"), Chunk("I live in Pune", "user")])

    def test_assistant_turn_becomes_reply(self):
        self.assertEqual(split_transcript("Alice: I am vegetarian\nBot: Noted."),
                         [Chunk("I am vegetarian", "user", "Noted.")])

    def test_consecutive_assistant_turns_concatenate(self):
        self.assertEqual(split_transcript("Alice: hi\nBot: hello\nAssistant: how are you"),
                         [Chunk("hi", "user", "hello how are you")])

    def test_leading_assistant_turn_dropped(self):
        self.assertEqual(split_transcript("Agent: welcome\nAlice: hi"), [Chunk("hi", "user")])

    def test_same_speaker_turns_merge_but_not_across_reply(self):
        chunks = split_transcript("Alice: hi\nAlice: I am vegetarian\nBot: ok\nAlice: I live in Pune")
        self.assertEqual(chunks, [Chunk("hi I am vegetarian", "user", "ok"),
                                  Chunk("I live in Pune", "user")])

    def test_continuation_line_appends(self):
        self.assertEqual(split_transcript("Alice: I am\nvegetarian now\nBot: ok"),
                         [Chunk("I am vegetarian now", "user", "ok")])

    def test_json_roles(self):
        text = json.dumps([{"role": "user", "content": "hi"}, {"role": "assistant", "content": "yo"}])
        self.assertEqual(split_transcript(text), [Chunk("hi", "user", "yo")])

    def test_json_assistant_role_wins_over_speaker_name(self):
        text = json.dumps([{"role": "user", "content": "hi"},
                           {"role": "assistant", "speaker": "Sam", "content": "yo"}])
        self.assertEqual(split_transcript(text), [Chunk("hi", "user", "yo")])

    def test_long_turn_cut_with_reply_on_last_piece(self):
        chunks = split_transcript(f"Alice: {sentences(42)}\nBot: got it")
        self.assertGreaterEqual(len(chunks), 2)
        self.assertTrue(all(len(c.text) <= CHUNK_CHARS for c in chunks))
        self.assertTrue(all(c.speaker == "user" for c in chunks))
        self.assertEqual([c.agent_reply for c in chunks], [None] * (len(chunks) - 1) + ["got it"])

    def test_unknown_owner_lists_speakers(self):
        with self.assertRaises(ValueError) as cm:
            split_transcript("Alice: hi\nBob: hey\nBot: hello", owner="Carol")
        self.assertIn("Alice", str(cm.exception))
        self.assertIn("Bob", str(cm.exception))


class PlanTest(unittest.TestCase):
    def test_returns_kind_and_chunks(self):
        self.assertEqual(plan_chunks("a.txt", b"Alice is vegetarian."),
                         ("prose", [Chunk("Alice is vegetarian.")]))
        kind, chunks = plan_chunks("", b"Alice: hi\nBot: hello")
        self.assertEqual((kind, chunks), ("transcript", [Chunk("hi", "user", "hello")]))

    def test_format_override(self):
        kind, chunks = plan_chunks("a.txt", b"Diet: veg\nCity: Pune\nJob: nurse", fmt="transcript")
        self.assertEqual(kind, "transcript")
        self.assertEqual([c.text for c in chunks], ["veg", "Pune", "nurse"])

    def test_bad_format(self):
        with self.assertRaisesRegex(ValueError, "format"):
            plan_chunks("a.txt", b"hello", fmt="poem")

    def test_too_many_bytes(self):
        with self.assertRaisesRegex(ValueError, "MB"):
            plan_chunks("a.txt", b"a" * (ingest.MAX_BYTES + 1))

    def test_too_many_chunks(self):
        old = ingest.MAX_CHUNKS
        ingest.MAX_CHUNKS = 2
        try:
            with self.assertRaisesRegex(ValueError, "limit of 2"):
                plan_chunks("a.txt", b"Alice: a\nBob: b\nAlice: c\nBob: d")
        finally:
            ingest.MAX_CHUNKS = old

    def test_json_not_turn_list(self):
        with self.assertRaisesRegex(ValueError, "json"):
            plan_chunks("a.json", b'{"name": "Alice"}')

    def test_empty_text(self):
        with self.assertRaisesRegex(ValueError, "no text"):
            plan_chunks("a.txt", b"  \n ")



class ReviewFixesTest(unittest.TestCase):
    """Regressions from the /memories review: detection, preamble, timestamps, JSON roles,
    mid-block headings, deep JSON and the DOCX size guards."""

    def test_short_key_value_notes_are_prose(self):
        for text in ("Allergies: peanuts\nBlood type: O+", "Allergy: peanuts",
                     "Q: where do you live?\nA: Pune",
                     "Subject: Dinner\nFrom: Bob\nHi, Alice is allergic to shellfish.",
                     "Meeting notes\nAction: Alice sends deck\nAction: Bob books room\n"
                     "Action: review\nDecision: go"):
            with self.subTest(text=text):
                self.assertEqual(detect_format(text), "prose")
        self.assertEqual(plan_chunks("", b"Allergy: peanuts"), ("prose", [Chunk("Allergy: peanuts")]))

    def test_back_and_forth_without_assistant_is_transcript(self):
        self.assertEqual(detect_format("Alice: hi\nBob: hello\nAlice: I move in May"), "transcript")

    def test_preamble_before_first_turn_is_kept(self):
        kind, chunks = plan_chunks(
            "", b"Call with Bob about the Pune move on Monday.\nAlice: hi\nBob: hello\nAlice: I move in May")
        self.assertEqual(kind, "transcript")
        self.assertEqual(chunks[0], Chunk("Call with Bob about the Pune move on Monday. hi", "user"))
        self.assertEqual([c.speaker for c in chunks], ["user", "Bob", "user"])

    def test_timestamped_chat_export_is_transcript(self):
        text = b"[10:31] Alice: I am vegetarian\n[10:32] Bob: I love steak\n[10:33] Alice: ok"
        self.assertEqual(plan_chunks("", text), ("transcript", [
            Chunk("I am vegetarian", "user"), Chunk("I love steak", "Bob"), Chunk("ok", "user")]))
        wa = "10/03/2024, 10:31 - Alice: hi\n10/03/2024, 10:32 - Bob: yo\n10/03/2024, 10:33 - Alice: ok"
        self.assertEqual([c.speaker for c in split_transcript(wa)], ["user", "Bob", "user"])

    def test_json_system_and_tool_roles_skipped(self):
        text = json.dumps([
            {"role": "system", "content": "You are a helpful assistant. The user likes cats."},
            {"role": "user", "content": [{"type": "text", "text": "I live"}, {"type": "text", "text": "in Pune"}]},
            {"role": "assistant", "content": None, "tool_calls": []},
            {"role": "tool", "content": "42"},
            {"role": "assistant", "content": "Nice"}])
        self.assertEqual(plan_chunks("chat.json", text.encode()),
                         ("transcript", [Chunk("I live in Pune", "user", "Nice")]))

    def test_heading_inside_block_starts_new_chunk(self):
        self.assertEqual([c.text for c in split_prose("# Diet\nShe is vegetarian.\n# Work\nShe is a nurse.")],
                         ["# Diet\nShe is vegetarian.", "# Work\nShe is a nurse."])

    def test_deeply_nested_json_is_value_error_not_crash(self):
        with self.assertRaises(ValueError):
            plan_chunks("a.json", b"[" * 100000)
        self.assertEqual(plan_chunks("", b"[" * 100000)[0], "prose")

    def test_docx_zip_bomb_rejected_before_parsing(self):
        buf = io.BytesIO()
        with zipfile.ZipFile(buf, "w", zipfile.ZIP_DEFLATED) as z:
            z.writestr("word/document.xml", b"<w:p/>" * (ingest.MAX_DOCX_XML // 6 + 1))
        with self.assertRaisesRegex(ValueError, "too large"):
            extract_text("bomb.docx", buf.getvalue())

    def test_docx_many_paragraphs_is_fast_and_capped(self):
        paras = [f"Fact number {i}." for i in range(5000)]
        data = tiny_docx(*paras)
        t0 = time.monotonic()
        text = extract_text("big.docx", data)
        self.assertLess(time.monotonic() - t0, 1.5)          # was ~2.6 s: p.style per paragraph
        self.assertTrue(text.startswith("# Notes\n\nFact number 0."))
        old = ingest.MAX_TEXT
        ingest.MAX_TEXT = 1000
        try:
            with self.assertRaisesRegex(ValueError, "Split it"):
                extract_text("big.docx", data)
        finally:
            ingest.MAX_TEXT = old


if __name__ == "__main__":
    unittest.main(verbosity=2)
